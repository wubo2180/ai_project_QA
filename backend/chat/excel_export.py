"""Excel 导出工具模块。"""
import io
import re

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def has_markdown_table(text: str) -> bool:
    """检测文本中是否包含 Markdown 表格。"""
    # 查找表格分隔行，如 |---|---|
    return bool(re.search(r"^\|[\s\-:|+]+\|$", text, re.MULTILINE))


def extract_tables(text: str) -> list[list[list[str]]]:
    """从 Markdown 文本中提取所有表格。
    
    返回: [[表1], [表2], ...]，每个表 = [表头行, 数据行1, 数据行2, ...]
    """
    lines = text.strip().split("\n")
    tables = []
    current_table = []
    in_table = False

    for line in lines:
        stripped = line.strip()
        # 跳过空行
        if not stripped:
            if in_table and current_table:
                tables.append(current_table)
                current_table = []
            in_table = False
            continue

        # 检查是否是表格行 (以 | 开头和结尾)
        if stripped.startswith("|") and stripped.endswith("|"):
            # 分隔行 (|---|) 跳过
            if re.match(r"^\|[\s\-:|+]+\|$", stripped):
                in_table = True
                continue
            # 解析单元格
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            current_table.append(cells)
            in_table = True
        else:
            if in_table and current_table:
                tables.append(current_table)
                current_table = []
            in_table = False

    if in_table and current_table:
        tables.append(current_table)

    return tables


def tables_to_excel(tables: list[list[list[str]]]) -> io.BytesIO:
    """将表格列表写入 Excel 文件，返回 BytesIO 对象。"""
    wb = openpyxl.Workbook()
    # 删除默认 sheet
    wb.remove(wb.active)

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    cell_font = Font(size=11)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, table in enumerate(tables):
        sheet_name = f"表格{idx + 1}" if len(tables) > 1 else "Sheet1"
        ws = wb.create_sheet(title=sheet_name[:31])  # Excel sheet 名最多 31 字符

        for row_idx, row_data in enumerate(table):
            for col_idx, cell_value in enumerate(row_data):
                cell = ws.cell(row=row_idx + 1, column=col_idx + 1, value=cell_value)
                cell.font = header_font if row_idx == 0 else cell_font
                cell.fill = header_fill if row_idx == 0 else PatternFill()
                cell.alignment = center_align
                cell.border = thin_border

        # 自动调整列宽
        for col_idx in range(len(table[0])):
            max_len = 0
            for row_data in table:
                if col_idx < len(row_data):
                    # 中文字符算 2 个宽度
                    cell_len = sum(2 if ord(c) > 127 else 1 for c in row_data[col_idx])
                    max_len = max(max_len, cell_len)
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx + 1)].width = min(max_len + 4, 60)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
